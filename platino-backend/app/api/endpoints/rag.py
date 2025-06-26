from fastapi import APIRouter, UploadFile, File, HTTPException, Depends
from sqlalchemy.orm import Session
from typing import List
from io import BytesIO
from llama_index.readers.file import PyMuPDFReader
from fastapi.middleware.cors import CORSMiddleware
from llama_index.readers.file import PDFReader
from llama_index.core.node_parser import SimpleNodeParser
from llama_index.core import Document as LlamaDocument
from docx import Document as DocxDocument
import tempfile

from ...db.session import get_db
from ...db.models import Document
from ...db.chunks import add_chunks

router = APIRouter()

# In memory storage of uploaded files (name only)
FILES: List[str] = []

@router.post("/files")
async def upload_file(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Upload a file to be used for retrieval."""
    try:
        content = await file.read()
        # placeholder: store content in real vector store
        FILES.append(file.filename)
        doc = Document(filename=file.filename)
        db.add(doc)
        db.commit()
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
    return {"filename": file.filename}

@router.get("/files")
async def list_files(db: Session = Depends(get_db)):
    """List uploaded files."""
    docs = db.query(Document).all()
    filenames = [doc.filename for doc in docs]
    return {"files": filenames}


@router.post("/split_pdf")
async def split_pdf(file: UploadFile = File(...)):
    """Return PDF chunks using LlamaIndex."""
    try:
        contents = await file.read()
        pdf_reader = PDFReader()
        documents = pdf_reader.load_data(BytesIO(contents))
        parser = SimpleNodeParser.from_defaults()
        nodes = parser.get_nodes_from_documents(documents)
        chunks = [node.text for node in nodes]
        return {"chunks": chunks}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/files_2")
async def upload_and_split_file(file: UploadFile = File(...)):
    """Upload a PDF or DOCX and return its chunks."""
    if not (file.filename.endswith(".pdf") or file.filename.endswith(".docx")):
        raise HTTPException(status_code=400, detail="Solo se permiten archivos PDF o DOCX")

    try:
        content = await file.read()

        documents = []
        if file.filename.endswith(".pdf"):
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            reader = PyMuPDFReader()
            documents = reader.load_data(file_path=tmp_path)
        else:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            doc = DocxDocument(tmp_path)
            text = "\n".join([p.text for p in doc.paragraphs])
            documents = [LlamaDocument(text=text)]

        parser = SimpleNodeParser()
        nodes = parser.get_nodes_from_documents(documents)
        chunks = [node.text for node in nodes]

        add_chunks(file.filename, chunks)

        return {"filename": file.filename, "chunks": chunks}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))